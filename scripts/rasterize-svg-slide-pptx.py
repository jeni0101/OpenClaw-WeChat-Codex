#!/usr/bin/env python3
"""Rebuild a one-SVG-per-slide PPTX as a standard JPEG-slide presentation."""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    parser.add_argument("--slides", type=int, default=27)
    parser.add_argument("--quality", type=int, default=88)
    args = parser.parse_args()
    if not 80 <= args.quality <= 95:
        parser.error("--quality must be between 80 and 95")

    source = args.source.resolve(strict=True)
    destination = args.destination.resolve(strict=False)
    if source == destination:
        raise SystemExit("destination must differ from source")
    destination.parent.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory(
        prefix=".pptx-raster-",
        dir=destination.parent,
    ) as temporary_name:
        temporary = Path(temporary_name)
        rendered: list[Path] = []
        with zipfile.ZipFile(source) as archive:
            names = set(archive.namelist())
            expected = [f"ppt/media/image{index}.svg" for index in range(1, args.slides + 1)]
            missing = [name for name in expected if name not in names]
            if missing:
                raise SystemExit(f"missing SVG slide artwork: {missing}")
            for index, name in enumerate(expected, start=1):
                svg = temporary / f"slide-{index:02d}.svg"
                jpg = temporary / f"slide-{index:02d}.jpg"
                svg.write_bytes(archive.read(name))
                result = subprocess.run(
                    [
                        "/usr/bin/convert",
                        "-background", "white",
                        "-alpha", "remove",
                        "-density", "144",
                        str(svg),
                        "-resize", "1600x900!",
                        "-strip",
                        "-interlace", "Plane",
                        "-quality", str(args.quality),
                        str(jpg),
                    ],
                    stdin=subprocess.DEVNULL,
                    stdout=subprocess.PIPE,
                    stderr=subprocess.PIPE,
                    timeout=120,
                    check=False,
                )
                if result.returncode != 0 or not jpg.is_file():
                    message = result.stderr.decode("utf-8", errors="replace")[-2000:]
                    raise SystemExit(f"failed to render slide {index}: {message}")
                rendered.append(jpg)

        presentation = Presentation()
        presentation.slide_width = Inches(13.333333)
        presentation.slide_height = Inches(7.5)
        blank = presentation.slide_layouts[6]
        for image in rendered:
            slide = presentation.slides.add_slide(blank)
            slide.shapes.add_picture(
                str(image),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )
        staging = temporary / "standard.pptx"
        presentation.save(staging)

        checked = Presentation(str(staging))
        if len(checked.slides) != args.slides:
            raise SystemExit(f"rebuilt slide count mismatch: {len(checked.slides)}")
        with zipfile.ZipFile(staging) as archive:
            if archive.testzip() is not None:
                raise SystemExit("rebuilt PPTX has a corrupt ZIP member")
            media = [
                name for name in archive.namelist()
                if name.startswith("ppt/media/") and not name.endswith("/")
            ]
            if len(media) != args.slides or not all(name.lower().endswith(".jpg") for name in media):
                raise SystemExit(f"unexpected rebuilt media set: {media}")
            external = [
                name for name in archive.namelist()
                if name.endswith(".rels") and b'TargetMode="External"' in archive.read(name)
            ]
            if external:
                raise SystemExit(f"rebuilt PPTX contains external relationships: {external}")
        shutil.copyfile(staging, destination)

    report = {
        "source": str(source),
        "destination": str(destination),
        "slides": args.slides,
        "jpegSlides": args.slides,
        "size": destination.stat().st_size,
        "sha256": sha256(destination),
    }
    print(json.dumps(report, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
