#!/usr/bin/env python3
"""Build deterministic, standards-compliant 17-slide PPTX size fixtures."""

from __future__ import annotations

import argparse
import hashlib
import json
import random
import sys
import tempfile
from pathlib import Path

VENDORED_PACKAGES = Path(__file__).resolve().parents[1] / "data" / "tools" / "python-pptx"
if VENDORED_PACKAGES.is_dir():
    sys.path.insert(0, str(VENDORED_PACKAGES))

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


TARGETS = [
    64 * 1024,
    128 * 1024,
    192 * 1024,
    256 * 1024,
    512 * 1024,
    1024 * 1024,
    1540 * 1024,
    2 * 1024 * 1024,
]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build(destination: Path, dimension: int, seed: int) -> int:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    with tempfile.TemporaryDirectory(prefix="pptx-matrix-") as temp_dir:
        temp = Path(temp_dir)
        for slide_index in range(17):
            rng = random.Random(seed + slide_index)
            image = Image.new("RGB", (dimension, dimension))
            image.putdata(
                [
                    (rng.randrange(256), rng.randrange(256), rng.randrange(256))
                    for _ in range(dimension * dimension)
                ]
            )
            image_path = temp / f"slide-{slide_index + 1:02d}.png"
            image.save(image_path, format="PNG", compress_level=9, optimize=False)
            slide = presentation.slides.add_slide(blank)
            slide.shapes.add_picture(
                str(image_path),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )
        presentation.save(destination)
    return destination.stat().st_size


def build_nearest(destination: Path, target: int) -> dict[str, int | str]:
    lower = 1
    upper = 1024
    best: tuple[int, int] | None = None
    for _ in range(13):
        dimension = (lower + upper) // 2
        size = build(destination, dimension, target)
        distance = abs(size - target)
        if best is None or distance < best[0]:
            best = (distance, dimension)
        if size < target:
            lower = dimension + 1
        else:
            upper = max(dimension - 1, 1)
        if lower > upper:
            break
    assert best is not None
    final_size = build(destination, best[1], target)
    return {
        "target": target,
        "size": final_size,
        "dimension": best[1],
        "sha256": sha256_file(destination),
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("output_dir", type=Path)
    parser.add_argument(
        "--targets-kib",
        help="comma-separated KiB targets; defaults to the historical 64 KiB–2 MiB matrix",
    )
    args = parser.parse_args()
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    targets = TARGETS
    if args.targets_kib:
        values = [value.strip() for value in args.targets_kib.split(",") if value.strip()]
        if not values:
            parser.error("--targets-kib must contain at least one integer")
        try:
            targets = [int(value) * 1024 for value in values]
        except ValueError:
            parser.error("--targets-kib must contain only integers")
        if any(target < 32 * 1024 or target > 50 * 1024 * 1024 for target in targets):
            parser.error("targets must be between 32 KiB and 50 MiB")
    reports = []
    for target in targets:
        label = f"{target // 1024}k"
        destination = output_dir / f"matrix-{label}.pptx"
        reports.append(build_nearest(destination, target))
    print(json.dumps(reports, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
