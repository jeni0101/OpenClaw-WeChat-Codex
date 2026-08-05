#!/usr/bin/env python3
"""Strict, read-only PPTX package and parser validation.

The validator intentionally does not repair, rewrite, compress, or rename the
source file. It prints a JSON report suitable for the media diagnostics bundle.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import sys
from pathlib import Path, PurePosixPath
from zipfile import BadZipFile, ZipFile

from lxml import etree
from pptx import Presentation


REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def relation_source(rels_name: str) -> str:
    rels = PurePosixPath(rels_name)
    if rels_name == "_rels/.rels":
        return ""
    if rels.parent.name != "_rels" or not rels.name.endswith(".rels"):
        raise ValueError(f"invalid relationship part path: {rels_name}")
    return str(rels.parent.parent / rels.name[:-5])


def resolve_target(source_name: str, target: str) -> str | None:
    if "://" in target or target.startswith("#"):
        return None
    base = posixpath.dirname(source_name)
    normalized = posixpath.normpath(posixpath.join(base, target))
    if normalized.startswith("../") or normalized == ".." or normalized.startswith("/"):
        raise ValueError(f"relationship escapes package: {source_name} -> {target}")
    return normalized


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    args = parser.parse_args()
    source = args.path.resolve(strict=True)
    report: dict[str, object] = {
        "path": str(source),
        "size": source.stat().st_size,
        "sha256": sha256_file(source),
        "errors": [],
    }
    errors: list[str] = report["errors"]  # type: ignore[assignment]

    try:
        with ZipFile(source) as archive:
            names = archive.namelist()
            name_set = set(names)
            duplicates = sorted({name for name in names if names.count(name) > 1})
            unsafe = sorted(
                name
                for name in names
                if name.startswith("/")
                or ".." in PurePosixPath(name).parts
                or "\\" in name
            )
            if duplicates:
                errors.append(f"duplicate package entries: {duplicates}")
            if unsafe:
                errors.append(f"unsafe package entries: {unsafe}")

            for required in (
                "[Content_Types].xml",
                "_rels/.rels",
                "ppt/presentation.xml",
                "ppt/_rels/presentation.xml.rels",
            ):
                if required not in name_set:
                    errors.append(f"missing required part: {required}")

            parser_xml = etree.XMLParser(
                resolve_entities=False,
                no_network=True,
                recover=False,
                huge_tree=False,
            )
            xml_parts = [name for name in names if name.endswith((".xml", ".rels"))]
            xml_roots: dict[str, etree._Element] = {}
            for name in xml_parts:
                try:
                    xml_roots[name] = etree.fromstring(archive.read(name), parser=parser_xml)
                except etree.XMLSyntaxError as exc:
                    errors.append(f"invalid XML {name}: {exc}")

            content_root = xml_roots.get("[Content_Types].xml")
            overrides: set[str] = set()
            defaults: set[str] = set()
            if content_root is not None:
                overrides = {
                    item.get("PartName", "").lstrip("/")
                    for item in content_root.findall(f"{{{CT_NS}}}Override")
                }
                defaults = {
                    item.get("Extension", "").lower()
                    for item in content_root.findall(f"{{{CT_NS}}}Default")
                }
                for name in names:
                    if name.endswith("/") or name == "[Content_Types].xml":
                        continue
                    extension = (
                        "rels"
                        if name.endswith(".rels")
                        else PurePosixPath(name).suffix.lstrip(".").lower()
                    )
                    if name not in overrides and extension not in defaults:
                        errors.append(f"missing content type for part: {name}")

            relationships = 0
            external_relationships = 0
            for rels_name, root in xml_roots.items():
                if not rels_name.endswith(".rels"):
                    continue
                source_name = relation_source(rels_name)
                for rel in root.findall(f"{{{REL_NS}}}Relationship"):
                    relationships += 1
                    target = rel.get("Target", "")
                    if rel.get("TargetMode") == "External":
                        external_relationships += 1
                        continue
                    try:
                        resolved = resolve_target(source_name, target)
                    except ValueError as exc:
                        errors.append(str(exc))
                        continue
                    if resolved and resolved not in name_set:
                        errors.append(f"missing relationship target: {rels_name} -> {resolved}")

            presentation_root = xml_roots.get("ppt/presentation.xml")
            xml_slide_count = 0
            if presentation_root is not None:
                xml_slide_count = len(
                    presentation_root.findall(f".//{{{P_NS}}}sldId")
                )

            image_entries = sorted(
                name
                for name in names
                if name.startswith("ppt/media/") and not name.endswith("/")
            )
            image_hashes = {
                name: hashlib.sha256(archive.read(name)).hexdigest()
                for name in image_entries
            }

            report.update(
                {
                    "packageEntries": len(names),
                    "xmlParts": len(xml_parts),
                    "relationships": relationships,
                    "externalRelationships": external_relationships,
                    "xmlSlideCount": xml_slide_count,
                    "imageCount": len(image_entries),
                    "imageSha256": image_hashes,
                }
            )
    except BadZipFile as exc:
        errors.append(f"invalid ZIP container: {exc}")

    try:
        presentation = Presentation(str(source))
        report["pythonPptxSlideCount"] = len(presentation.slides)
        report["pythonPptxShapeCounts"] = [
            len(slide.shapes) for slide in presentation.slides
        ]
    except Exception as exc:  # python-pptx surfaces several package exceptions
        errors.append(f"python-pptx rejected package: {type(exc).__name__}: {exc}")

    report["valid"] = not errors
    print(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True))
    return 0 if not errors else 1


if __name__ == "__main__":
    sys.exit(main())
