#!/usr/bin/env python3
"""Rebuild an image-slide PPTX with python-pptx without altering image bytes."""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import tempfile
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.util import Inches


def digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    parser.add_argument("--expected-slides", type=int, default=17)
    args = parser.parse_args()

    source = args.source.resolve(strict=True)
    destination = args.destination.resolve(strict=False)
    destination.parent.mkdir(parents=True, exist_ok=True)

    with ZipFile(source) as archive:
        image_names = [
            f"ppt/media/image{index}.png"
            for index in range(1, args.expected_slides + 1)
        ]
        missing = [name for name in image_names if name not in archive.namelist()]
        if missing:
            raise SystemExit(f"missing source images: {missing}")
        source_images = [(name, archive.read(name)) for name in image_names]

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank_layout = presentation.slide_layouts[6]

    with tempfile.TemporaryDirectory(prefix="pptx-standard-") as temp_dir:
        temp = Path(temp_dir)
        for index, (_, image_data) in enumerate(source_images, start=1):
            image_path = temp / f"slide-{index:02d}.png"
            image_path.write_bytes(image_data)
            slide = presentation.slides.add_slide(blank_layout)
            slide.shapes.add_picture(
                str(image_path),
                0,
                0,
                width=presentation.slide_width,
                height=presentation.slide_height,
            )
        staging = temp / "standard.pptx"
        presentation.save(staging)
        shutil.copyfile(staging, destination)

    rebuilt = Presentation(str(destination))
    if len(rebuilt.slides) != args.expected_slides:
        raise SystemExit(
            f"rebuilt slide count mismatch: {len(rebuilt.slides)}"
        )

    with ZipFile(destination) as archive:
        rebuilt_images = sorted(
            (
                name,
                archive.read(name),
            )
            for name in archive.namelist()
            if name.startswith("ppt/media/") and not name.endswith("/")
        )

    source_hashes = sorted(digest(data) for _, data in source_images)
    rebuilt_hashes = sorted(digest(data) for _, data in rebuilt_images)
    if source_hashes != rebuilt_hashes:
        raise SystemExit("image bytes changed during standards rebuild")

    report = {
        "source": str(source),
        "destination": str(destination),
        "slides": len(rebuilt.slides),
        "images": len(rebuilt_images),
        "imageBytesPreserved": True,
        "sourceImageSha256": source_hashes,
        "destinationSize": destination.stat().st_size,
        "destinationSha256": digest(destination.read_bytes()),
    }
    print(json.dumps(report, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
